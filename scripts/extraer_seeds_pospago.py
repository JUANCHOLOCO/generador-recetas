"""Extrae la parametrización inicial del módulo Pospago desde los Excel de la macro.

Se ejecuta UNA vez en desarrollo. Los CSV generados se revisan, se versionan en git y la
app los carga en su primera ejecución. La app nunca lee Excel.

Regla de combinación: un dato presente en ambos libros queda como AMBAS; presente solo en
uno, con el tipo de facturación de ese libro.

Uso:
    python scripts/extraer_seeds_pospago.py --anticipada "<ruta>.xlsm" --corriente "<ruta>.xlsm"
"""

from __future__ import annotations

import argparse
import csv
from collections.abc import Sequence
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
from openpyxl.worksheet.worksheet import Worksheet

from recetas.core.domain.billing import BillingScope
from recetas.modules.creacion_planes_pospago.domain.tipos import (
    ListaCambioPlan,
    TablaScoring,
    TipoPaquete,
)

HOJA = "Reglas Aplicabilidad"
PRIMERA_FILA = 3  # filas 1 y 2 son títulos
DESTINO = Path(__file__).resolve().parents[1] / "src/recetas/modules/creacion_planes_pospago/seeds"

# La celda J14 dice "CONTROL", pero la columna real de PLANS es TYPE (su valor es 'CONTROL').
CORRECCIONES_PLANS = {"CONTROL": "TYPE"}
# PLANS.BILLING_TYPE no se guarda: lo aporta la estrategia de facturación (Fase 4).
OMITIDAS = {("PLANS", "BILLING_TYPE")}

COLUMNAS_CSV: dict[str, tuple[str, ...]] = {
    "cpl_plan_plan.csv": ("id_plan", "lista"),
    "scoring.csv": ("tabla_destino", "scoring_id"),
    "canal.csv": ("nombre",),
    "campo_estatico.csv": ("tabla", "tipo_paquete", "columna", "valor_sql"),
    "cadena_pool.csv": ("part_num",),
    "plan_53_old.csv": ("id_plan", "part_num"),
    "rol.csv": ("role_granted",),
    "addon.csv": ("part_num",),
}

# Valores fijos de la familia Nintendo (columna AQ): plan=21, SMS=22, datos=23.
IDENTIFY = (("PLAN", "21"), ("SMS", "22"), ("DATA", "23"))

Fila = tuple[str, ...]


def texto(valor: object) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def columna(hoja: Worksheet, letra: str) -> list[str]:
    """Valores no vacíos de una columna, desde PRIMERA_FILA."""
    indice = column_index_from_string(letra)
    filas = hoja.iter_rows(min_row=PRIMERA_FILA, min_col=indice, max_col=indice, values_only=True)
    return [t for (valor,) in filas if (t := texto(valor))]


def pares(hoja: Worksheet, letra_clave: str, letra_valor: str) -> list[tuple[str, str]]:
    """Pares (clave, valor) de dos columnas contiguas, omitiendo filas sin clave."""
    desde = column_index_from_string(letra_clave)
    hasta = column_index_from_string(letra_valor)
    resultado = []
    for fila in hoja.iter_rows(
        min_row=PRIMERA_FILA, min_col=desde, max_col=hasta, values_only=True
    ):
        clave, valor = texto(fila[0]), texto(fila[-1])
        if clave:
            resultado.append((clave, valor))
    return resultado


def extraer(hoja: Worksheet) -> dict[str, list[Fila]]:
    cpl: list[Fila] = []
    # Mismo orden en que las listas aparecen en el script: Prepago, Nintendo, ATP, Control.
    for letra, lista in (
        ("B", ListaCambioPlan.PREPAGO),
        ("A", ListaCambioPlan.NINTENDO),
        ("D", ListaCambioPlan.ATP),
        ("C", ListaCambioPlan.CONTROL),
    ):
        cpl += [(id_plan, lista.value) for id_plan in columna(hoja, letra)]

    scoring = [(TablaScoring.SCORING_PLAN.value, v) for v in columna(hoja, "F")]
    scoring += [(TablaScoring.CPL_SCORING_PLAN.value, v) for v in columna(hoja, "G")]

    campos: list[Fila] = []
    for tabla, tipo, letra_columna, letra_valor in (
        ("PLANS", TipoPaquete.PLAN, "J", "K"),
        ("SERVICES", TipoPaquete.DATA, "M", "N"),
        ("SERVICES", TipoPaquete.SMS, "P", "Q"),
        ("PLAN_SERVICE", TipoPaquete.DATA, "S", "T"),
        ("PLAN_SERVICE", TipoPaquete.SMS, "V", "W"),
    ):
        for nombre, valor in pares(hoja, letra_columna, letra_valor):
            if tabla == "PLANS":
                nombre = CORRECCIONES_PLANS.get(nombre, nombre)
            if (tabla, nombre) not in OMITIDAS:
                campos.append((tabla, tipo.value, nombre, valor))

    return {
        "cpl_plan_plan.csv": cpl,
        "scoring.csv": scoring,
        "canal.csv": [(v,) for v in columna(hoja, "H")],
        "campo_estatico.csv": campos,
        "cadena_pool.csv": [(v,) for v in columna(hoja, "AB")],
        "plan_53_old.csv": pares(hoja, "AD", "AE"),
        "rol.csv": [(v,) for v in columna(hoja, "AS")],
        "addon.csv": [(v,) for v in columna(hoja, "AU")],
    }


def extraer_libro(ruta: Path) -> dict[str, list[Fila]]:
    # read_only: lectura en streaming, mucho más rápida en libros grandes.
    # data_only: valores calculados de las fórmulas, no la fórmula.
    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        return extraer(libro[HOJA])
    finally:
        libro.close()


def combinar(anticipada: list[Fila], corriente: list[Fila]) -> list[tuple[Fila, BillingScope]]:
    en_anticipada, en_corriente = set(anticipada), set(corriente)
    resultado = [
        (fila, BillingScope.AMBAS if fila in en_corriente else BillingScope.ANTICIPADA)
        for fila in anticipada
    ]
    resultado += [(fila, BillingScope.CORRIENTE) for fila in corriente if fila not in en_anticipada]
    return resultado


def escribir(ruta: Path, columnas: Sequence[str], filas: list[tuple[Fila, BillingScope]]) -> None:
    with ruta.open("w", encoding="utf-8", newline="") as archivo:
        writer = csv.writer(archivo, lineterminator="\n")
        writer.writerow(["orden", *columnas, "tipo_facturacion"])
        for orden, (fila, scope) in enumerate(filas, start=1):
            writer.writerow([orden, *fila, scope.value])


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("--anticipada", type=Path, required=True)
    parser.add_argument("--corriente", type=Path, required=True)
    parser.add_argument("--destino", type=Path, default=DESTINO)
    args = parser.parse_args(argv)

    anticipada = extraer_libro(args.anticipada)
    corriente = extraer_libro(args.corriente)
    args.destino.mkdir(parents=True, exist_ok=True)

    for nombre, columnas in COLUMNAS_CSV.items():
        filas = combinar(anticipada[nombre], corriente[nombre])
        escribir(args.destino / nombre, columnas, filas)
        print(f"{nombre:<22}{len(filas):>5} filas")

    identify = [(fila, BillingScope.AMBAS) for fila in IDENTIFY]
    escribir(args.destino / "identify.csv", ("tipo_paquete", "identify"), identify)
    print(f"{'identify.csv':<22}{len(identify):>5} filas")


if __name__ == "__main__":
    raise SystemExit(main())
